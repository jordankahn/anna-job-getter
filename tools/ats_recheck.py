import re, collections
from docx import Document
import docx2txt

JOBS = {
"Chicago Shakespeare": (
 "/home/user/anna-job-getter/jobs/chicago-shakespeare-theater-donor-services-associate/Anna_Brown_Resume_chicago-shakespeare-theater.md",
 """Donor Services Associate. accurate gift processing, stewardship, data integrity. Manage contributions in Tessitura, constituent data, donor communications and recognition, reporting and analysis, fundraising and engagement. Partner to Development, Ticketing, and Finance. Process gifts, pledges, and payments in Tessitura. tax receipts and acknowledgment letters. pledge schedules. Reconcile gifts with Finance. donor concierge ticketing account assistance. donor recognition listings. donor profiles giving histories. donor and prospect records. Data hygiene duplicate resolution. campaign appeal fund codes in Tessitura. Segmentation list building data extraction direct mail email gala. reports. donor retention growth campaign performance. Liaison Development Ticketing. Administrative support Chief Development Officer scheduling note-taking. Event support Gala. Required: Bachelor's degree, relational database system, Microsoft Office Excel Word, fundraising ticketing or nonprofit operations. Preferred: performing arts ticketing marketing donor stewardship reporting campaign tracking."""),
"Obama Foundation": (
 "/home/user/anna-job-getter/jobs/obama-foundation-senior-associate-principal-gifts/Anna_Brown_Resume_obama-foundation.md",
 """Senior Associate Principal Gifts Engagement. Support Vice President principal giving strategy program. operational pipelines, prospect briefing, metrics. donor portfolio tracking pipeline. Executive liaison calendar fundraising actions. fundraising strategies. briefing solicitations meetings. proposals stewardship materials briefing memos. liaison Donor Relations Events. Prospect Research. deliverables Salesforce trackers. data integrity portfolios tracking metrics. Required: 3-4 years nonprofit development, strategic operations, or fundraising support. sensitive donor information discretion integrity. project management multi-layered processes cross-functional teams. written oral communication board-ready briefings stewardship collateral. systems literacy donor management Salesforce data analytics. Hybrid Chicago."""),
}
STOP = set("a an the and or of to in for with on at by from as is are be this that role you your we our their will all across into within during other more most can able team teams work working experience including high strong support provide ensure manage managing maintain new program programs help may etc end day days time type position company employee applicant person status year years professional minimum based".split())
SECTIONS = ["SUMMARY","SKILLS","EXPERIENCE","EDUCATION","AWARDS"]

def load(p):
    blocks=[]
    for raw in open(p):
        s=raw.rstrip("\n"); t=s.strip()
        if not t: continue
        if t.startswith("# "): blocks.append(("name",t[2:])); continue
        if t.startswith("## "): blocks.append(("sec",t[3:].upper())); continue
        t=t.replace("**","")
        if t.startswith("- "): blocks.append(("body",t[2:])); continue
        if (s.startswith("  ") or s.startswith("\t")) and blocks:
            blocks[-1]=(blocks[-1][0], blocks[-1][1]+" "+t); continue
        blocks.append(("body",t))
    return blocks

def build_docx(blocks,path):
    d=Document()
    for typ,txt in blocks:
        if typ=="name":
            p=d.add_paragraph(); r=p.add_run(txt); r.bold=True
        elif typ=="sec": d.add_heading(txt,level=1)
        else: d.add_paragraph(txt)
    d.save(path)

def keywords(text,n=28):
    words=re.findall(r"[a-zA-Z][a-zA-Z\-]{3,}",text.lower())
    freq=collections.Counter(w for w in words if w not in STOP)
    phrases=["gift processing","data integrity","tessitura","salesforce","stewardship","pledge","acknowledgment","segmentation","reconcile","ticketing","donor","database","fundraising","reporting","nonprofit","excel","project management","operations","briefing","portfolio","pipeline","discretion","metrics","cross-functional","communication","records","campaign","finance","liaison"]
    out=[]
    for k in [w for w,_ in freq.most_common(n)]+phrases:
        if k in text.lower() and k not in out: out.append(k)
    return out

print("="*68)
print("CORRECTED (honest) resumes — re-test")
print("="*68)
for name,(mdpath,jd) in JOBS.items():
    blocks=load(mdpath); dp=f"/tmp/{name.replace(' ','_')}_corrected.docx"; build_docx(blocks,dp)
    txt=docx2txt.process(dp); low=txt.lower()
    email=bool(re.search(r"[\w.]+@[\w.]+",txt)); phone=bool(re.search(r"\d{3}[.\-]\d{3}[.\-]\d{4}",txt))
    secs=[s for s in SECTIONS if s.lower() in low]
    kws=keywords(jd); found=[k for k in kws if k.lower() in low]; score=round(100*len(found)/len(kws))
    print(f"\n### {name}")
    print(f"  Clean extraction: {'YES' if len(txt)>300 else 'NO'} | Email:{'Y' if email else 'N'} Phone:{'Y' if phone else 'N'} | Sections {len(secs)}/5")
    print(f"  Honest keyword match: {score}%  ({len(found)}/{len(kws)})")
print("\n"+"="*68)
