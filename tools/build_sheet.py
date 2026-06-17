import openpyxl, base64
from openpyxl.styles import Font, PatternFill, Alignment

wb = openpyxl.Workbook()

bold = Font(bold=True)
title_font = Font(bold=True, size=14)
hdr_font = Font(bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="4A148C")
sec_font = Font(bold=True, size=11, color="4A148C")
note_align = Alignment(wrap_text=True, vertical="top")
grey = Font(italic=True, color="888888")

# ---------- Tab 1: Start Here ----------
ws1 = wb.active
ws1.title = "Start Here"
ws1.sheet_properties.tabColor = "4A148C"
rows1 = [
    ("Anna Job Tracker", title_font),
    ("", None),
    ("How this works", bold),
    ("1. Paste any job link into the 'My Job Links' tab (one per row). Add the date and any notes.", None),
    ("2. Tell Claude 'check the sheet'. Claude reads your links, checks each against the filters, and builds tailored materials.", None),
    ("3. For each job, Claude creates a folder in Drive: Anna Job Search > Applications > <Org> - <Role>, with a Resume, a Cover Letter (only if the posting needs one), and an Evaluation.", None),
    ("4. Status and links are logged in the project's TRACKER.md (in GitHub) and in the Applications folder.", None),
    ("", None),
    ("Jobs Claude finds for you appear in the 'Claude-Sourced Jobs' tab (only roles posted within the last 7 days).", None),
    ("", None),
    ("Every job must pass these filters:", bold),
    ("   - Salary at least $50,000/year (listed salary preferred)", None),
    ("   - Chicagoland (in-person or hybrid) OR fully remote", None),
    ("   - Full-time", None),
    ("", None),
    ("Note: Claude can create and read this sheet but cannot edit its cells automatically. It will not tick boxes here; it writes results to the Applications folder and TRACKER.md and reports the links to you.", None),
    ("", None),
    ("Applications folder: https://drive.google.com/drive/folders/1SU6qQ0iKWfMMMz0-c5WLinZUXGvxIAig", None),
]
for i, (text, font) in enumerate(rows1, 1):
    c = ws1.cell(row=i, column=1, value=text)
    if font: c.font = font
    c.alignment = note_align
ws1.column_dimensions["A"].width = 110

# ---------- Tab 2: My Job Links ----------
ws2 = wb.create_sheet("My Job Links")
ws2.sheet_properties.tabColor = "1B5E20"
headers2 = ["Date Added", "Job URL", "Title (optional)", "Cover Letter? (if known)", "Notes"]
for j, h in enumerate(headers2, 1):
    c = ws2.cell(row=1, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill
example = ["2026-06-17", "https://example.org/jobs/123  (example row - delete me)", "Program Coordinator", "", "Paste real job links below, one per row"]
for j, v in enumerate(example, 1):
    c = ws2.cell(row=2, column=j, value=v); c.font = grey
widths2 = [14, 60, 26, 24, 40]
for j, w in enumerate(widths2, 1):
    ws2.column_dimensions[chr(64+j)].width = w
ws2.freeze_panes = "A2"

# ---------- Tab 3: Claude-Sourced Jobs ----------
ws3 = wb.create_sheet("Claude-Sourced Jobs")
ws3.sheet_properties.tabColor = "B71C1C"
cols3 = ["Date Posted", "Title", "Organization", "Location", "Salary", "Fit Grade", "Source Board", "Job URL", "Status"]
widths3 = [13, 34, 30, 18, 20, 9, 14, 70, 26]
for j, w in enumerate(widths3, 1):
    ws3.column_dimensions[chr(64+j)].width = w

r = 1
ws3.cell(row=r, column=1, value="Claude-Sourced Jobs  -  only roles posted within the last 7 days").font = title_font; r += 1
ws3.cell(row=r, column=1, value="Last search: 2026-06-17").font = grey; r += 2

ws3.cell(row=r, column=1, value="Posted within the last 7 days (June 10-17, 2026)").font = sec_font; r += 1
for j, h in enumerate(cols3, 1):
    c = ws3.cell(row=r, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill
r += 1
note = ("None verified this run. Live job boards block automated date-sorted access (HTTP 403), and mid-June is a "
        "slow posting week, so no role could be confirmed as posted June 10-17. Use the bookmarked 'sort by newest' "
        "board links at the bottom of this tab in a browser to catch this-week posts.")
c = ws3.cell(row=r, column=1, value=note); c.alignment = note_align; c.font = Font(italic=True)
ws3.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
ws3.row_dimensions[r].height = 45
r += 2

ws3.cell(row=r, column=1, value="Recent strong matches OUTSIDE the 7-day window  (NOT sourced per your filter - verify the post date and that the role is still open before applying)").font = sec_font
ws3.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9); r += 1
for j, h in enumerate(cols3, 1):
    c = ws3.cell(row=r, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill
r += 1

data = [
    ["May 7, 2026", "Program Associate (arts/EJ grantmaking)", "Prince Charitable Trusts", "Chicago (hybrid)", "$76,500-92,000", "A", "Idealist", "https://www.idealist.org/en/job/ee621bda000f4260aa043c76ce2bbfee-program-associate-prince-charitable-trusts-chicago", "Verify open (deadline ~June 5 may have passed)"],
    ["Apr 15, 2026", "Director of Programs", "Chicago Cultural Alliance", "Chicago (hybrid)", "$70,000", "A", "Idealist", "https://www.idealist.org/en/nonprofit-job/0c7ae7f5e6a24de0a6bf1539adcfaa96-director-of-programs-chicago-cultural-alliance-chicago", "Verify open"],
    ["Mar 19, 2026", "Grants Associate", "Paul M. Angell Family Foundation", "Chicago (hybrid)", "$80,000-90,000", "A-", "Idealist", "https://www.idealist.org/en/nonprofit-job/47834024ddf3480e9364ee24ed20ed54-grants-associate-paul-m-angell-family-foundation-pmaff-chicago", "Verify open"],
    ["Feb 24, 2026", "Associate, Development", "Chicago Public Education Fund", "Chicago (hybrid)", "$59,000-74,000", "B+", "Idealist", "https://www.idealist.org/en/nonprofit-job/f007677eedfc4a1d883585f707456479-associate-development-the-chicago-public-education-fund-chicago", "Verify open"],
    ["Dec 15, 2025", "Development Director", "Chicago Jazz Philharmonic", "Chicago (hybrid)", "$50,000+", "A-", "Idealist", "https://www.idealist.org/en/nonprofit-job/80d8e0ea9ae24280876c949c8cac0fea-development-director-chicago-jazz-philharmonic-chicago", "Verify open"],
    ["early 2026", "Senior Director of Development", "Arts Alliance Illinois", "IL (90% remote)", "$102,000-113,500", "B", "Idealist", "https://www.idealist.org/en/nonprofit-job/d06d966b0b404de3ab6f5d697beeebdb-senior-director-of-development-arts-alliance-illinois-chicago", "Senior stretch; verify open"],
    ["Oct 29, 2025", "Community Arts Coordinator", "National Museum of Mexican Art", "Chicago", "$48,000-53,000", "B+", "Idealist", "https://www.idealist.org/en/nonprofit-job/b583ed93af374c43b61d6dc6f0ee04e5-community-arts-coordinator-national-museum-of-mexican-art-chicago", "Salary partly under $50k; likely filled"],
    ["Sep 18, 2025", "Development & Communications Associate", "Chicago Arts Partnerships in Education (CAPE)", "Chicago (hybrid)", "$45,000-65,000", "A", "Idealist", "https://www.idealist.org/en/nonprofit-job/94d6544b67fe45a9b69e0754951a7f6a-development-communications-associate-chicago-arts-partnerships-in-education-chicago", "Salary partly under $50k; likely filled"],
]
for row in data:
    for j, v in enumerate(row, 1):
        ws3.cell(row=r, column=j, value=v).alignment = note_align
    r += 1
r += 1

ws3.cell(row=r, column=1, value="Bookmark these 'sort by newest' boards to catch this-week posts (they block automation but work in a browser):").font = bold
r += 1
boards = [
    "League of Chicago Theatres jobs: https://jobs.leagueofchicagotheatres.org/",
    "Idealist Chicago arts/music (sort: newest): https://www.idealist.org/en/arts-music-jobs-chicago-il",
    "Americans for the Arts Job Bank (Illinois): https://jobbank.artsusa.org/l-illinois-jobs.html",
    "Chicago Cultural Alliance opportunities: https://www.chicagoculturalalliance.org/opportunities/",
    "Foundation List Chicago: https://www.foundationlist.org/chicago-nonprofit-jobs/",
    "ArtSearch/TCG administration jobs: https://artjobs.artsearch.us/job-listings/administration-jobs/",
]
for b in boards:
    ws3.cell(row=r, column=1, value=b); r += 1
ws3.freeze_panes = "A1"

wb.save("/tmp/anna_tracker.xlsx")
with open("/tmp/anna_tracker.xlsx", "rb") as f:
    b64 = base64.b64encode(f.read()).decode()
with open("/tmp/b64.txt", "w") as f:
    f.write(b64)
print("bytes:", len(b64))
