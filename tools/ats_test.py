import re, os, collections
from docx import Document
from docx.shared import Pt
import docx2txt

RESUMES = {
"Chicago Shakespeare - Donor Services Associate": """Anna Brown
Chicago, IL 60640 | 847-395-5749 | anna.brown5749@gmail.com | www.annaleebrown.com

SUMMARY
Detail-oriented administrative and donor services professional with a Master's in Arts Leadership and Cultural Management and hands-on experience managing CRM records, memberships, and constituent relationships for a Chicago cultural nonprofit. Skilled in data integrity and accuracy, gift and payment processing, fundraising operations, donor and member communication, reporting, and event support, with a genuine commitment to the arts. Comfortable in Salesforce and other relational database systems and known for careful follow-up and warm, professional service.

SKILLS
Donor services and stewardship support | Relational database management (Salesforce; quick to learn Tessitura) | Processing gifts, pledges, and payments | Acknowledgments and donor recognition | Data hygiene, data integrity, and records management | Campaign and appeal coding and reporting | Reporting, list building, and segmentation | Ticketing and patron services (Eventive) | Membership management | Donor and member communications | Fundraising and development operations support | Gala and donor event support | Customer service | Invoicing and reconciliation with finance (Stampli) | Microsoft Office (Excel, Word)

EXPERIENCE
JCC Chicago - Chicago, IL
Operations Coordinator, Community Engagement | Oct 2025 - Present
Manage administrative operations and accurate records for the Community Engagement team using internal systems.
Code and submit invoicing and reconcile financial transactions with the finance team, maintaining data integrity and accuracy.
Co-direct the annual Jewish Chicago Film Festival, building and managing ticketing in Eventive and supporting event logistics, guest coordination, and community-partner and donor relationships.
Develop and deliver programs and events for families, adults, and seniors.
Senior Customer Engagement Representative | Aug 2024 - Oct 2025
Managed membership records, POS and payment transactions, and constituent communications, maintaining accurate data and running list segmentation in Salesforce and Upace.
Built and maintained relationships that supported retention and a positive member experience, the foundation of strong donor stewardship.

EDUCATION
Colorado State University (Online) - M.A., Arts Leadership and Cultural Management; Graduate Certificate, Nonprofit Administration
Carthage College - Kenosha, WI - B.A., Music Theatre; Summa Cum Laude (3.9 GPA)

AWARDS
Chicago Chapter NATS Music Theatre Competition, 3rd Place; Kennedy Center American College Theatre Festival, Irene Ryan Semi-Finalist (2020)
""",
"Obama Foundation - Senior Associate Principal Gifts": """Anna Brown
Chicago, IL 60640 | 847-395-5749 | anna.brown5749@gmail.com | www.annaleebrown.com

SUMMARY
Strategic operations and program coordinator with a Master's in Arts Leadership and Cultural Management and experience running administrative operations, managing Salesforce data, and coordinating complex events and cross-functional projects for a mission-driven nonprofit. Detail-oriented and discreet, with strong project management, communication, and stewardship-minded relationship skills.

SKILLS
Strategic and program operations | Project and process management | Salesforce / CRM data integrity | Pipeline and portfolio tracking | Reporting, metrics, and data analytics | Cross-functional coordination and liaison | Briefing and materials preparation | Donor and stakeholder relations | Development and fundraising operations support | Event and logistics coordination | Executive and administrative support | Discretion with sensitive information | Written and verbal communication | Microsoft Office (Excel, Word, PowerPoint), Google Workspace

EXPERIENCE
JCC Chicago - Chicago, IL
Operations Coordinator, Community Engagement | Oct 2025 - Present
Run all administrative operations for the Community Engagement team, building systems, schedules, trackers, and pipelines that keep multi-layered projects on time and provide clear tracking metrics.
Manage data and records in Salesforce, maintaining accuracy and integrity across constituent information.
Co-direct the annual Jewish Chicago Film Festival, coordinating logistics, partners, and stakeholders across teams and preparing materials and communications.
Serve as a liaison across departments to deliver programs and events for diverse audiences and align cross-functional deliverables.
Senior Customer Engagement Representative | Aug 2024 - Oct 2025
Maintained constituent records and sensitive member information with discretion in Salesforce.
Coordinated schedules and communications and supported relationship-building and stewardship.

EDUCATION
Colorado State University (Online) - M.A., Arts Leadership and Cultural Management; Graduate Certificate, Nonprofit Administration
Carthage College - Kenosha, WI - B.A., Music Theatre; Summa Cum Laude (3.9 GPA)

AWARDS
Chicago Chapter NATS Music Theatre Competition, 3rd Place; Kennedy Center American College Theatre Festival, Irene Ryan Semi-Finalist (2020)
""",
}

JDS = {
"Chicago Shakespeare - Donor Services Associate": """Donor Services Associate. Ensure exceptional donor experience through accurate gift processing, thoughtful stewardship, and strong data integrity. Manage all contributions in Tessitura, maintain high-quality constituent data, support donor communications and recognition, provide timely reporting and analysis to advance fundraising and engagement. Partner to Development, Ticketing, and Finance. Process and record all gifts, pledges, and payments in Tessitura. Generate personalized tax receipts and acknowledgment letters. Track pledge schedules. Reconcile gifts and pledges with Finance. Serve as donor concierge including ticketing and account assistance. Maintain donor recognition listings. Prepare donor profiles and giving histories. Maintain accurate donor and prospect records. Data hygiene, duplicate resolution, quality control. Set up campaign, appeal, fund, designation, and source codes in Tessitura. Segmentation, list building, and data extraction for direct mail, email, gala, tele-funding. Generate standard and ad-hoc reports. Analyze donor retention, donor growth, campaign performance. Liaison between Development and Ticketing. Administrative support to Chief Development Officer including scheduling and note-taking. Event support for Opening Nights, Gala. Required: Bachelor's degree, experience with a relational database system, high proficiency in Microsoft Office especially Excel and Word, experience in fundraising, ticketing, or nonprofit operations. Preferred: performing arts, ticketing, marketing, donor stewardship, reporting, campaign tracking.""",
"Obama Foundation - Senior Associate Principal Gifts": """Senior Associate, Principal Gifts and Engagement. Support the Vice President in launching and executing the principal giving strategy program. Manage operational pipelines, prospect briefing processes, and metrics. Manage the Vice President's premier donor portfolio, tracking movement, engagement schedules, next steps in the pipeline. Executive liaison, calendar alignment, track fundraising actions. Foundation-wide collaboration, unify fundraising strategies. Oversee briefing process for solicitations and meetings. Coordinate personalized proposals, stewardship materials, and briefing memos. Operational liaison with Donor Relations and Events. Collaborate with Prospect Research. Tracking deliverables across Salesforce and shared trackers. Maintain data integrity for portfolios, provide tracking metrics. Required: 3-4 years professional experience in nonprofit development, strategic operations, or executive-level frontline fundraising support. Handle sensitive donor information with discretion and integrity. Excellent project management skills, driving multi-layered processes across cross-functional teams. Excellent written and oral communication, board-ready briefings, stewardship collateral. Strong systems literacy, donor management systems Salesforce preferred, data analytics. Hybrid role based in Chicago.""",
}

STOP = set("a an the and or of to in for with on at by from as is are be this that role you your we our their will all across into within during other more most can able team teams work working experience including high strong support provide ensure manage managing maintain new program programs help may etc end day days time type position company employee applicant person status year years professional minimum based".split())
SECTIONS = ["SUMMARY","SKILLS","EXPERIENCE","EDUCATION","AWARDS"]

def keywords(text, n=28):
    words = re.findall(r"[a-zA-Z][a-zA-Z\-]{3,}", text.lower())
    freq = collections.Counter(w for w in words if w not in STOP)
    phrases = ["gift processing","data integrity","tessitura","salesforce","stewardship","pledge","acknowledgment","segmentation","reconcile","reconciliation","ticketing","donor","database","fundraising","reporting","nonprofit","excel","project management","operations","briefing","portfolio","pipeline","discretion","metrics","cross-functional","data analytics","communication","records","campaign","finance","liaison"]
    top = [w for w,_ in freq.most_common(n)]
    out=[]
    for k in top+phrases:
        if k in text.lower() and k not in out:
            out.append(k)
    return out

def build_docx(name, text, path):
    doc = Document()
    style = doc.styles["Normal"]; style.font.name="Calibri"; style.font.size=Pt(11)
    lines=[l.rstrip() for l in text.strip().split("\n")]
    for i,l in enumerate(lines):
        if not l: continue
        if i==0:
            h=doc.add_paragraph(); r=h.add_run(l); r.bold=True; r.font.size=Pt(20); continue
        if l.strip() in SECTIONS:
            doc.add_heading(l.strip(), level=1); continue
        if (" | " not in l) and l[0].isupper() and l.endswith(".") and len(l)>40:
            doc.add_paragraph(l, style="List Bullet"); continue
        doc.add_paragraph(l)
    doc.save(path)

print("="*70)
for name in RESUMES:
    docx_path=f"/tmp/{name.split(' - ')[0].replace(' ','_')}.docx"
    build_docx(name, RESUMES[name], docx_path)
    extracted = docx2txt.process(docx_path)
    low = extracted.lower()
    d = Document(docx_path)
    email = bool(re.search(r"[\w.]+@[\w.]+", extracted))
    phone = bool(re.search(r"\d{3}[.\-]\d{3}[.\-]\d{4}", extracted))
    secs_found=[s for s in SECTIONS if s.lower() in low]
    kws=keywords(JDS[name])
    found=[k for k in kws if k.lower() in low]
    missing=[k for k in kws if k.lower() not in low]
    score=round(100*len(found)/len(kws))
    print(f"\n### {name}")
    print(f"File: {os.path.basename(docx_path)}  ({len(extracted.split())} words)")
    print(f"  Clean text extraction: {'YES' if len(extracted)>400 else 'NO'} | Email: {'YES' if email else 'NO'} | Phone: {'YES' if phone else 'NO'}")
    print(f"  Sections {len(secs_found)}/5: {', '.join(secs_found)} | Tables: {len(d.tables)} | Images: {len(d.inline_shapes)} | Columns: 1")
    print(f"  KEYWORD MATCH: {score}%  ({len(found)}/{len(kws)})")
    print(f"  Still missing: {', '.join(missing) if missing else '(none)'}")
print("\n"+"="*70)
